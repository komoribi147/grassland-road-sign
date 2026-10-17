"""Build the readable production script from the published story JSON."""

import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
STORY = ROOT / "story"
OUTPUT = ROOT / "docs" / "草地上的路标完整互动剧本修订版.docx"


def load(rel):
    return json.loads((STORY / rel).read_text(encoding="utf-8"))


manifest = load("manifest.json")
characters = load("characters.json")
lore = load("lore.json")
scenes = [json.loads((ROOT / rel).read_text(encoding="utf-8")) for rel in manifest["scenes"]]

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.76)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)


def font(style, size, bold=False, color="000000"):
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


normal = doc.styles["Normal"]
font(normal, 10.3)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.22

for name, size, bold, before, after, color in [
    ("Title", 20, True, 0, 10, "000000"),
    ("Heading 1", 14.5, True, 16, 7, "000000"),
    ("Heading 2", 11.5, True, 11, 5, "000000"),
]:
    s = doc.styles[name]
    font(s, size, bold, color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

title_style_pr = doc.styles["Title"].element.get_or_add_pPr()
for border in title_style_pr.findall(qn("w:pBdr")):
    title_style_pr.remove(border)

for name, size, indent, before, after, color in [
    ("Script Visual", 9.7, 0.12, 0, 4, "474747"),
    ("Script Dialogue", 10.3, 0.16, 0, 4, "111111"),
    ("Script Option", 9.7, 0.38, 0, 3, "333333"),
    ("Script Cue", 9.2, 0.12, 0, 4, "666666"),
    ("Script Branch", 10.2, 0, 10, 5, "000000"),
    ("Script Reference", 9.6, 0, 0, 1, "333333"),
]:
    s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = normal
    font(s, size, name == "Script Branch", color)
    s.paragraph_format.left_indent = Inches(indent)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_together = True
    if name == "Script Branch":
        s.paragraph_format.keep_with_next = True


def para(text, style=None, prefix=None):
    p = doc.add_paragraph(style=style)
    if prefix:
        r = p.add_run(prefix)
        r.bold = True
    p.add_run(text.replace("{player}", "（当前称呼）"))
    return p


def title(text, style):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def display_name(who):
    return "玩家" if who == "__player__" else characters.get(who, {}).get("display", who)


def add_hyperlink(paragraph, label, url):
    part = paragraph.part
    relationship = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1B4D75")
    props.append(color)
    under = OxmlElement("w:u")
    under.set(qn("w:val"), "single")
    props.append(under)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = label
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def branch_caption(scene_id, name):
    labels = {
        ("act1", "observe"): "分支 A 走近老班长",
        ("act1", "skip"): "分支 B 留在原地",
        ("act1", "hejie"): "共同段",
        ("act3", "water"): "分支 A 走近何大姐",
        ("act3", "after_water"): "共同段",
        ("act3", "daughter"): "高信任补充对话",
        ("act3", "after_daughter"): "共同段",
    }
    return labels.get((scene_id, name), f"分支 {name}")


def emit_steps(steps, scene_id, nested=False):
    for step in steps:
        kind = step["t"]
        if kind == "visual":
            para(step["description"].replace("\n", "\n"), "Script Visual", "画面  ")
        elif kind == "say":
            para(step["text"], "Script Dialogue", display_name(step["who"]) + "  ")
        elif kind == "action":
            detail = step.get("hint", "")
            para(step["label"] + (f"（{detail}）" if detail else ""), "Script Cue", "玩家动作  ")
        elif kind == "choose":
            para(step.get("prompt", "请选择动作"), "Script Cue", "动作选择  ")
            for index, option in enumerate(step["options"], start=1):
                line = f"{index}. {option['text']}"
                if option.get("reply"):
                    line += f"  →  {display_name(option.get('replyWho', step.get('replyWho', '')))}：{option['reply']}"
                if option.get("goto"):
                    line += f"  →  {branch_caption(scene_id, option['goto'].split(':')[-1])}"
                para(line, "Script Option")
                emit_steps(option.get("after", []), scene_id, nested=True)
        elif kind == "select":
            para("玩家可选一句回答；也可自由表达。", "Script Cue", "玩家回应  ")
            for index, option in enumerate(step["options"], start=1):
                para(f"{index}. 玩家：{option['text']}  →  {display_name(step['ai'])}：{option['reply']}", "Script Option")
                emit_steps(option.get("after", []), scene_id, nested=True)
            fallback = step.get("fallback")
            if fallback:
                value = fallback[0] if isinstance(fallback[0], str) else fallback[0].get("text", "")
                para(f"自由表达离线接话：{display_name(step['ai'])}：{value}", "Script Cue")
        elif kind == "label":
            para(branch_caption(scene_id, step["name"]), "Script Branch")
        elif kind == "if":
            target = step.get("goto", "").split(":")[-1]
            para(f"信任达到 {step['cond']['value']} 时，进入“{branch_caption(scene_id, target)}”；其他情况直接进入下一个共同段。", "Script Cue", "分支条件  ")
        elif kind == "lore":
            para(lore[step["card"]]["title"], "Script Cue", "解锁史料卡  ")
        elif kind == "end":
            para(step["text"], "Script Dialogue", "片尾文字  ")
        if kind in ("visual", "say", "action", "select") and step.get("after"):
            emit_steps(step["after"], scene_id, nested=True)


title("草地上的路标 完整互动剧本修订版", "Title")
para("1935年8月下旬，红军右路军经松潘草地向班佑行进。玩家从泥水里醒来，跟随老班长李长河、司号员石头与护理员何秀英寻找前队。故事由序章、四幕和尾声组成，按数日行军中的关键片段展开。")
para("本剧人物及其具体遭遇为艺术创作；毛儿盖、草地、湘江、班佑等历史地点与事件依据公开史料。玩家的犹豫、害怕和请求帮助都属于正常反应，不设唯一的“正确情感”。")

title("人物", "Heading 1")
for key in ("li-changhe", "shitou", "he-xiuying"):
    person = characters[key]
    para(f"{person['display']}  {person['identity']} {person['persona']}")

title("贯穿全剧的线索", "Heading 1")
for item in (
    "草鞋：序章老班长把右脚草鞋给玩家；泥沼中他的右脚受伤；后段左脚草鞋也断了；结尾玩家仍穿着他给的那一只。",
    "路标：泥沼边的第一根木棍留给后队；石头后来把军号上的另一段红布系在老班长探路棍上。结尾插在老班长身旁的是第二根。",
    "名字：玩家认出伤员田德福的名字，石头因此提出“文书”的称呼；夜里又记住李长河、桂香与何秀英。",
    "时间：从毛儿盖附近进入草地，到看见班佑河是数日行军。黎明只是接近目的地，众人还须继续前行。",
):
    para(item)

script_heading = title("完整剧本", "Heading 1")
script_heading.paragraph_format.page_break_before = True
for scene in scenes:
    title(scene["title"], "Heading 2")
    if scene.get("ambient"):
        para("环境音以雨声为底，随场景强弱变化。", "Script Cue")
    emit_steps(scene["steps"], scene["id"])

title("史实与创作说明", "Heading 1")
para("剧中的人物姓名、失散亲人、救援过程、红布路标和晚年经历是虚构情节。史料卡将能核实的时代背景与这些创作情节分开表述。泥沼救援和寻找野外食物只是剧情，不构成现实操作指引。")
for item in lore.values():
    p = doc.add_paragraph(style="Script Reference")
    r = p.add_run(item["title"] + "  ")
    r.bold = True
    p.add_run(item["text"])
    p2 = doc.add_paragraph(style="Script Reference")
    p2.paragraph_format.left_indent = Inches(0.12)
    p2.add_run("史料参考  ")
    add_hyperlink(p2, item.get("sourceTitle", "资料链接"), item.get("sourceUrl", ""))

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("草地上的路标")
run.font.size = Pt(8)

doc.save(OUTPUT)
print(OUTPUT)
